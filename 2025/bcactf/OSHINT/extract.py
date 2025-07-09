from pptx import Presentation
from hashlib import sha256

prs = Presentation('hi.pptx')
hsh = '1db4650f72af7fc27414993e845ff0eceff18bb93c6a4381eb56efe1580ec9bb'
for i, slide in enumerate(prs.slides, 1):
    print(f"Slide {i}:")
    for shape in slide.shapes:
        # print(shape)
        # picture
        if shape.shape_type == 13:
            if shape.image:
                image = shape.image
                image_bytes = image.blob
                image_hash = sha256(image_bytes).hexdigest()
                # print(f"Image {i}: {image_hash}")
                if image_hash != hsh:
                    print(f"Image {i} hash does not match: {image_hash}")

        # text
        elif shape.has_text_frame:
            text = shape.text_frame.text
            if text:
                print(f"Text: {text}")